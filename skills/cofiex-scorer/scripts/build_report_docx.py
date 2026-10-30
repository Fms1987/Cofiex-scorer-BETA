"""Gera o relatório final da simulação COFIEX em .docx a partir de um JSON.

Uso: python build_report_docx.py relatorio.json saida.docx [--carta-cliente carta.docx]

O JSON segue o esquema descrito em references/output/docx-report.md. Os subtotais e a Nota
Geral são recalculados aqui a partir dos itens (com o teto da Nota de Área Estratégica), para
que as tabelas do sumário executivo nunca fiquem inconsistentes com os itens.

Com --carta-cliente, grava também a carta de comentários ao mutuário (bloco "carta_cliente"),
sem notas. Antes de salvar, avisa quando dois parágrafos de seções diferentes são quase iguais
(regra "um achado, uma seção").
"""
import difflib
import json
import sys
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HEADER_FILL = "1F3864"
SUBTOTAL_FILL = "D9E2F3"
TOTAL_FILL = "B4C6E7"
NIVEL_PONTOS = {"A": 5.0, "B": 3.0, "C": 2.0, "D": 1.0, "E": 0.5, "0": 0.0, "-": 0.0}


def fmt(v, casas=2):
    if v is None or v == "":
        return "—"
    if isinstance(v, str):
        return v
    return f"{v:.{casas}f}".replace(".", ",")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table(doc, header, rows, widths=None, styles=None, num_cols=()):
    """styles: lista paralela a rows com None | 'subtotal' | 'total'."""
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, HEADER_FILL)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        st = styles[r_i] if styles else None
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val))
            if i in num_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if st:
                run.bold = True
                shade(cells[i], TOTAL_FILL if st == "total" else SUBTOTAL_FILL)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def pts(area, k):
    """Pontos da área no cenário k (min/esp/max): valor numérico explícito, se houver, ou o nível."""
    v = area.get(f"pontos_{k}")
    return float(v) if v is not None else NIVEL_PONTOS.get(area[f"nivel_{k}"], 0.0)


def declaradas(areas):
    """Áreas que entram na soma: as declaradas no pleito (inclusive sem sustentação), não as sugeridas."""
    return [a for a in areas if a.get("pleiteado_como", "Declarada").startswith("Declarada")]


def soma(itens, chave):
    return sum((it.get(chave) or 0) for it in itens)


def subtotal_ae(d):
    """Nota de Área Estratégica por cenário: áreas declaradas + bônus, limitada ao teto."""
    ae, teto_ae = d["area_estrategica"], d.get("teto_area_estrategica", 10.0)
    areas = declaradas(ae["areas"])
    return {k: min(teto_ae, sum(pts(a, k) for a in areas) + soma(ae.get("bonus", []), k))
            for k in ("min", "esp", "max")}


def totais(d):
    sub_do = {k: 0 for k in ("min", "esp", "max")} if d.get("ente_federal", False) else \
        {k: soma(d["dados_objetivos"], k) for k in ("min", "esp", "max")}
    sub_ae = subtotal_ae(d)
    return {k: sub_do[k] + sub_ae[k] for k in sub_ae}, sub_ae


def tabela_pontuacao(doc, d):
    """Tabela 1: mínimo/esperado/máximo por item, subtotais e Nota Geral."""
    federal = d.get("ente_federal", False)
    teto_ae = d.get("teto_area_estrategica", 10.0)
    rows, styles = [], []
    sub_do = {"min": 0, "esp": 0, "max": 0}
    if not federal:
        for it in d["dados_objetivos"]:
            rows.append(["Dados Objetivos", it["item"], fmt(it["min"]), fmt(it["esp"]), fmt(it["max"])])
            styles.append(None)
        sub_do = {k: soma(d["dados_objetivos"], k) for k in ("min", "esp", "max")}
        rows.append(["Subtotal", "Dados Objetivos (teto 5,00)", fmt(sub_do["min"]), fmt(sub_do["esp"]), fmt(sub_do["max"])])
        styles.append("subtotal")

    ae = d["area_estrategica"]
    areas = declaradas(ae["areas"])
    for a in areas:
        rows.append(["Área Estratégica", f"Área {a['codigo']}",
                     fmt(pts(a, "min")), fmt(pts(a, "esp")), fmt(pts(a, "max"))])
        styles.append(None)
    for b in ae.get("bonus", []):
        situacao = f" ({b['situacao']})" if b.get("situacao") else ""
        rows.append(["Área Estratégica", f"Bônus — {b['item']}{situacao}", fmt(b["min"]), fmt(b["esp"]), fmt(b["max"])])
        styles.append(None)
    sub_ae = subtotal_ae(d)
    rows.append(["Subtotal", f"Nota de Área Estratégica (teto {fmt(teto_ae)}, bônus incluído)",
                 fmt(sub_ae["min"]), fmt(sub_ae["esp"]), fmt(sub_ae["max"])])
    styles.append("subtotal")

    total = {k: sub_do[k] + sub_ae[k] for k in ("min", "esp", "max")}
    rotulo = "Nota Geral (= Nota de Área Estratégica, ente federal)" if federal else \
        "Nota Geral (Dados Objetivos + Nota de Área Estratégica, teto 15,00)"
    rows.append(["Total", rotulo, fmt(total["min"]), fmt(total["esp"]), fmt(total["max"])])
    styles.append("total")
    table(doc, ["Bloco", "Item", "Mínimo", "Esperado", "Máximo"], rows,
          widths=[3.0, 8.0, 2.0, 2.0, 2.0], styles=styles, num_cols=(2, 3, 4))
    return total


def tabela_areas(doc, d):
    """Tabela 2: objetivos estratégicos do PPA indicados no pleito e pontuação possível."""
    rows = []
    for a in d["area_estrategica"]["areas"]:
        faixa = f"{a['nivel_min']} / {a['nivel_esp']} / {a['nivel_max']}"
        pts = f"{fmt(NIVEL_PONTOS.get(a['nivel_min'], 0))} / {fmt(NIVEL_PONTOS.get(a['nivel_esp'], 0))} / " \
              f"{fmt(NIVEL_PONTOS.get(a['nivel_max'], 0))}"
        rows.append([a["codigo"], a["objetivo"], a.get("pleiteado_como", "Declarada"), faixa, pts,
                     a.get("justificativa_curta", "")])
    table(doc, ["Código PPA", "Objetivo estratégico", "Situação", "Nível (mín./esp./máx.)",
                "Pontos (mín./esp./máx.)", "Fundamento resumido"], rows,
          widths=[1.6, 4.4, 2.0, 2.2, 2.6, 4.6], num_cols=(0, 3, 4))
    if any(a not in declaradas(d["area_estrategica"]["areas"]) for a in d["area_estrategica"]["areas"]):
        bloco(doc, {"tipo": "nota", "texto": "Objetivos marcados como sugeridos não foram declarados no pleito "
                                             "e não entram no subtotal da Tabela 1."})


def tabela_sugestoes(doc, d):
    """Tabela 3: sugestões de melhoria, resumidas (o detalhe vai no Plano de Correção)."""
    rows = [[str(i + 1), s["sugestao"], s.get("item_afetado", ""), s.get("prioridade", ""),
             s.get("ganho_estimado", "")] for i, s in enumerate(d.get("sugestoes", []))]
    table(doc, ["#", "Sugestão", "Item afetado", "Prioridade", "Ganho estimado"], rows,
          widths=[0.8, 8.2, 3.2, 2.0, 2.6], num_cols=(0, 3, 4))


def bloco(doc, b):
    tipo = b.get("tipo", "p")
    if tipo == "p":
        doc.add_paragraph(b["texto"])
    elif tipo == "bullets":
        for item in b["itens"]:
            doc.add_paragraph(item, style="List Bullet")
    elif tipo == "numerada":
        for item in b["itens"]:
            doc.add_paragraph(item, style="List Number")
    elif tipo == "tabela":
        if b.get("titulo"):
            doc.add_paragraph().add_run(b["titulo"]).bold = True
        table(doc, b["cabecalho"], b["linhas"])
    elif tipo == "nota":
        p = doc.add_paragraph()
        run = p.add_run(b["texto"])
        run.italic = True
        run.font.size = Pt(9)


def novo_documento():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
    return doc


def nota_destaque(doc, rotulo, valor, tamanho, idioma="pt"):
    p = doc.add_paragraph()
    texto = fmt(valor) if idioma == "pt" else f"{valor:.2f}"
    r = p.add_run(f"{rotulo}: {texto}")
    r.bold = True
    r.font.size = Pt(tamanho)


ROTULOS_TTL = {
    "pt": {"titulo": "Resumo para o líder de equipe", "veredito": "Veredito",
           "geral": "Nota Geral esperada", "ae": "Nota de Área Estratégica esperada",
           "sublimite": "Sublimite recomendado", "lacunas": "As três lacunas que mais custam nota",
           "acoes": "As três ações de maior ganho", "ganho": "ganho"},
    "en": {"titulo": "Summary for the task team leader", "veredito": "Verdict",
           "geral": "Expected Overall Score", "ae": "Expected Strategic Area Score",
           "sublimite": "Recommended sublimit", "lacunas": "The three gaps that cost the most points",
           "acoes": "The three highest-gain actions", "ganho": "gain"},
}


def resumo_ttl(doc, d, total, sub_ae, idioma):
    """Resumo do líder de equipe (1–2 páginas): um número em destaque, o resto em segundo plano."""
    r = d["resumo_ttl"] if idioma == "pt" else d["resumo_ttl"]["en"]
    L = ROTULOS_TTL[idioma]
    doc.add_heading(L["titulo"], level=1)
    p = doc.add_paragraph()
    p.add_run(f"{L['veredito']}: ").bold = True
    p.add_run(r["veredito"])
    nota_destaque(doc, L["geral"], total["esp"], 16, idioma)
    nota_destaque(doc, L["ae"], sub_ae["esp"], 12, idioma)
    if r.get("frase_faixa"):
        bloco(doc, {"tipo": "nota", "texto": r["frase_faixa"]})
    p = doc.add_paragraph()
    p.add_run(f"{L['sublimite']}: ").bold = True
    p.add_run(r["sublimite_recomendado"])
    doc.add_paragraph().add_run(L["lacunas"]).bold = True
    # Numeração manual: a lista numerada do Word continuaria a contagem entre as duas listas.
    for i, t in enumerate(r["lacunas"][:3], start=1):
        doc.add_paragraph(f"{i}. {t}")
    doc.add_paragraph().add_run(L["acoes"]).bold = True
    for i, a in enumerate(r["acoes"][:3], start=1):
        doc.add_paragraph(f"{i}. {a['acao']} ({L['ganho']}: {a['ganho']})")


def carta_cliente(d, out):
    """Carta de comentários ao mutuário, por seção do Anexo I. Sem notas nem leitura de sublimite."""
    c = d["carta_cliente"]
    doc = novo_documento()
    doc.add_heading("Comentários à carta-consulta", level=0)
    doc.add_paragraph(f"{d['cabecalho']['proposta']} — {d['cabecalho']['proponente']}").runs[0].bold = True
    if c.get("introducao"):
        doc.add_paragraph(c["introducao"])
    rows = [[it["secao_anexo"], it["problema"], it["redacao_sugerida"]] for it in c["itens"]]
    table(doc, ["Seção do formulário (Anexo I)", "Problema", "Redação sugerida"], rows, widths=[3.5, 5.5, 8.0])
    if c.get("fechamento"):
        doc.add_paragraph(c["fechamento"])
    doc.save(out)
    print(f"Carta ao cliente salva em {out} — {len(rows)} comentários")


def checar_repeticao(d, limite=0.85, minimo=80):
    """Alerta parágrafos quase iguais em seções diferentes (cada achado deve aparecer uma vez)."""
    textos = []
    for sec in d.get("secoes", []):
        for b in sec.get("blocos", []):
            for t in ([b.get("texto")] if b.get("texto") else b.get("itens", [])):
                if t and len(t) >= minimo:
                    textos.append((sec["titulo"], t))
    alertas = []
    for i, (s1, t1) in enumerate(textos):
        for s2, t2 in textos[i + 1:]:
            if s1 != s2 and difflib.SequenceMatcher(None, t1, t2).ratio() >= limite:
                alertas.append(f"- \"{s1}\" e \"{s2}\": \"{t1[:70]}…\"")
    if alertas:
        print("AVISO: parágrafos quase iguais em seções diferentes. Manter na seção dona e "
              "trocar o outro por uma remissão:", *alertas, sep="\n", file=sys.stderr)
    return alertas


def main(src, out, out_cliente=None):
    d = json.loads(open(src, encoding="utf-8").read())
    checar_repeticao(d)
    doc = novo_documento()

    cab = d["cabecalho"]
    doc.add_heading("Simulação da Avaliação Técnica COFIEX", level=0)
    doc.add_paragraph(f"{cab['proposta']} — {cab['proponente']}").runs[0].bold = True
    versao = f" Versão da skill: {d['versao_skill']}." if d.get("versao_skill") else ""
    doc.add_paragraph(f"Data da simulação: {d.get('data', date.today().strftime('%d/%m/%Y'))}.{versao} "
                      "Simulação sob julgamento cego, sem consulta à Nota Técnica do pleito. "
                      "Resultado estimativo; não substitui a avaliação oficial.").runs[0].italic = True

    total_calc, sub_ae_calc = totais(d)
    if d.get("resumo_ttl"):
        resumo_ttl(doc, d, total_calc, sub_ae_calc, "pt")
        if d["resumo_ttl"].get("en"):
            resumo_ttl(doc, d, total_calc, sub_ae_calc, "en")
        doc.add_page_break()

    doc.add_heading("Sumário executivo", level=1)
    table(doc, ["Campo", "Valor"], [["Proposta", cab["proposta"]], ["Proponente", cab["proponente"]]] +
          [list(c) for c in cab.get("campos", [])], widths=[5.0, 12.0])
    if d.get("sintese"):
        doc.add_paragraph(d["sintese"])

    doc.add_heading("Tabela 1 — Pontuação por item: mínima, esperada e máxima", level=2)
    total = tabela_pontuacao(doc, d)
    doc.add_paragraph(f"Nota Geral esperada: {fmt(total['esp'])}.").runs[0].bold = True
    fe = d.get("faixa_empirica")
    if fe:
        bloco(doc, {"tipo": "nota", "texto": f"Faixa calibrada pelos erros do banco de teste (cobre cerca de 80% "
                                             f"dos casos): Nota Geral de {fmt(fe['min'])} a {fmt(fe['max'])}. "
                                             "As colunas Mínimo e Máximo acima são cenários da rubrica."})
    if d.get("base_calibracao"):
        bloco(doc, {"tipo": "nota", "texto": "Base de calibração: " + d["base_calibracao"]})

    doc.add_heading("Tabela 2 — Objetivos estratégicos do PPA indicados no pleito e pontuação possível", level=2)
    tabela_areas(doc, d)

    doc.add_heading("Tabela 3 — Sugestões a considerar (resumo)", level=2)
    tabela_sugestoes(doc, d)
    bloco(doc, {"tipo": "nota", "texto": "Sugestões indicativas. O detalhamento, a priorização e o ganho "
                                         "estimado de cada uma estão na seção Plano de Correção."})

    for sec in d.get("secoes", []):
        doc.add_heading(sec["titulo"], level=1)
        for b in sec.get("blocos", []):
            bloco(doc, b)

    doc.save(out)
    print(f"Relatório salvo em {out} — Nota Geral esperada {fmt(total['esp'])}")
    if out_cliente:
        if d.get("carta_cliente"):
            carta_cliente(d, out_cliente)
        else:
            print("AVISO: --carta-cliente pedido, mas o JSON não tem o bloco carta_cliente.", file=sys.stderr)


if __name__ == "__main__":
    a = sys.argv[1:]
    cliente = a[a.index("--carta-cliente") + 1] if "--carta-cliente" in a else None
    main(a[0], a[1], cliente)
