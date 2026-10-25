"""Build docs/analise-areas-estrategicas.md from validated area records + ficha headers.

usage: build_areas_doc.py <final_dir> <fichas_root> <out_md>
"""
import json, re, sys
from collections import Counter, defaultdict
from pathlib import Path

LVL = {"A": 5.0, "B": 3.0, "C": 2.0, "D": 1.0, "E": 0.5}


def header(p):
    return dict(re.findall(r"^(\w+): ?(.*)$", p.read_text(encoding="utf-8").split("---", 2)[1], re.M))


def fmt(x, nd=3):
    try:
        return f"{float(x):.{nd}f}".rstrip("0").rstrip(".") if float(x) != int(float(x)) else f"{int(float(x))}"
    except (TypeError, ValueError):
        return str(x)


SECTION1 = """# Análise das áreas estratégicas nos pleitos da Cofiex — esquema e registro por pleito

> Documento de referência, separado do skill. Os pleitos são identificados por códigos opacos (`P…` pleito, `A…` avaliação), os
> mesmos das fichas em `fichas/`. Rótulos de áreas, justificativas e decisões estão em inglês (padrão do corpus). Fonte: Notas
> Técnicas (NTs) da SEAID; as pontuações finais divulgadas pela Cofiex foram conferidas na planilha de avaliações.

## 1. Esquema de análise das áreas estratégicas

### 1.1 Resolução Cofiex nº 17/2021 (regime anterior) — "análise técnica" em dois blocos

A NT pontua cada pleito em **A) Análise técnica (0–4,0)** e **B) Áreas estratégicas (0–2,0)**; o IDH entra à parte no total.

| Bloco | Peso | Itens (máximo) |
|---|---|---|
| **A.1 Desenho** | 1,8 | 1.1 justificativa do programa/projeto (0,7) · 1.2 objetivos, componentes e produtos alinhados à solução (0,6) · 1.3 coerência entre indicadores e resultados esperados (0,3) · 1.4 prazo de desembolso compatível e distribuição adequada (0,1) · 1.5 justificativa da opção pelo financiamento externo e pelo agente financeiro (0,1) |
| **A.2 Gestão/implementação** | 1,0 | 2.1 e 2.2 (somam 0,7, por diferença; inclui sustentabilidade e custos de manutenção, com nota de "qualidade do requisito") · 2.3 compatibilidade com as prioridades das políticas do proponente (0,3) |
| **A.3 Risco** | 1,2 | 3.1 (0,4, por diferença) · 3.2 identificação de riscos e ações mitigadoras (0,4) · 3.3 experiência na execução com o agente financeiro (0,4; usa o desempenho médio da carteira) |
| **B. Áreas estratégicas** | 2,0 | 15 critérios, tudo ou nada: 14 de 0,133 e o critério 1.15 de 0,138 (soma 2,0), com justificativa mesmo quando a nota é zero |

Os 15 critérios: 1.1 sustentabilidade ambiental, adaptação e mitigação · 1.2 eficiência energética e energias renováveis · 1.3 saneamento básico ·
1.4 gestão integrada de recursos hídricos e segurança hídrica · 1.5 biodiversidade e recursos florestais · 1.6 ambiente de negócios ·
1.7 empreendedores individuais e micro/pequenas empresas · 1.8 emprego e renda para populações vulneráveis ·
1.9 infraestrutura e equipamentos sociais para populações vulneráveis · 1.10 inclusão de minorias sociais · 1.11 juventude, gênero e raça ·
1.12 prestação de serviços públicos, gestão e participação cidadã · 1.13 educação e saúde · 1.14 desenvolvimento, inovação e modernização
tecnológica · 1.15 arranjos de financiamento que incentivem a participação privada.

- **Pontos adicionais (bônus):** o regime de 2021 não tem mecanismo de pontos adicionais.
- **Observado no corpus:** entre as NTs com os 15 critérios legíveis, os pleitos pontuaram entre 1 e 8 critérios (a maioria entre 3 e 5).

### 1.2 Resolução Normativa nº 1/2024 (regime atual) — admissibilidade e adequação às áreas estratégicas

1. **Análise técnica = admissibilidade.** Verifica os requisitos do Anexo I do PFE (resultado "Aprovação" ou não). Já não gera nota; é um
   portão. Reprovado nesse portão, o pleito não é pontuado.
2. **Adequação às áreas estratégicas (0–10).** O proponente indica na seção II.1 do pleito os objetivos do **Anexo II** (35 objetivos do PPA
   2024–2027, códigos 101–307) a que o programa atende. A NT avalia **somente os objetivos indicados**, dá a cada um **0 a 5 pontos** por nível
   e soma até 10:

   | Nível | A | B | C | D | E | não aceito |
   |---|---|---|---|---|---|---|
   | Pontos | 5,0 | 3,0 | 2,0 | 1,0 | 0,5 | 0 |

   O nível depende de: relação direta do objetivo geral com a área; indicadores SMART com linha de base; parcela de recursos aplicada à área;
   viabilidade do cronograma; e, para o nível A, caráter estrutural/inovador.
3. **Pontos adicionais (art. 15):** **§3º** +1,0 se os recursos forem integralmente direcionados a uma ou mais prioridades da Administração
   Pública Federal no PPA (a NT exige que o proponente o declare); **§4º** +1,0 pela promoção da integração regional (p.ex. zona de fronteira).
   O teto do AE continua em 10.
4. **Componentes do total (subnacionais):** AE (máx. 10) + IDH (máx. 2) + CAPAG (máx. 2) + trajetória da dívida (máx. 1) = 15.
   Federais: somente AE.
- **Observado no corpus (NTs de 2024 legíveis):** mapa nível→pontos confirmado (A=5, B=3, C=2, D=1, E=0,5); a maioria dos pleitos pontua 3 a 6
  objetivos; o ponto adicional do §3º foi concedido em cerca de 1 de cada 6 avaliações e o do §4º em casos raros.

### 1.3 Diferenças que importam para a avaliação
| Aspecto | 2021 | 2024 |
|---|---|---|
| Unidade de pontuação | 15 critérios fixos (0,133; o 1.15 vale 0,138) | até 35 objetivos do PPA (0–5 por nível) |
| Teto das áreas | 2,0 | 10 |
| Análise técnica | 4,0 pontos (desenho, gestão, risco) | admissibilidade (sem nota) |
| Bônus | não há | §3º (PPA) e §4º (integração regional), +1,0 cada |
| Declaração do proponente | informa o enquadramento; a NT justifica todos os 15 critérios | indica os objetivos; só eles são avaliados |

## 2. Registro por pleito
Cada avaliação lista: (a) áreas pontuadas com pontos e nível; (b) áreas apresentadas pelo proponente ou avaliadas que **não foram aceitas ou
pontuadas**, com uma justificativa curta; (c) pontos adicionais (pedidos? concedidos?) e a justificativa. NTs sem bloco de pontuação por área
(pautas, notas de outro tipo) aparecem como "sem pontuação por área na NT".
"""


def render_eval(r, h):
    L = []
    reg = "Res. Normativa 1/2024" if r["regime"] == "NR1_2024" else "Res. Cofiex 17/2021"
    meta = f"**{r['id']}** · tentativa {h.get('attempt', '?')} · {h.get('year', '?')}-T{h.get('quarter', '?')} · {reg} · desfecho `{h.get('outcome', '?')}`"
    if h.get("total"):
        meta += f" · total (planilha) {h['total']}"
    L.append(f"###### {meta}")
    if r["status"] != "ok" or not (r.get("areas") or r.get("described_unscored")):
        L.append("_Sem pontuação por área na NT (a NT disponível é pauta ou nota sem bloco de áreas)._\n")
        return L
    at = r.get("areas_total")
    if at:
        L.append(f"Áreas estratégicas: **{fmt(at[0])}/{fmt(at[1])}**" + (f" · AE na planilha: {h['ae']}" if r["regime"] == "NR1_2024" and h.get("ae") else ""))
    if r.get("described_unscored"):
        L.append("\n**Critérios descritos pela NT como atendidos (nota por critério não impressa; só o total acima)**\n")
        for a in r["described_unscored"]:
            L.append(f"- {a['code']} {a['label']}: {a.get('reason','')}")
    sc = [a for a in r["areas"] if float(a["points"]) > 0]
    ns = [a for a in r["areas"] if float(a["points"]) <= 0]
    if sc:
        L.append("\n**Pontuadas**\n\n| Código | Área | Nível | Pontos | Justificativa |\n|---|---|---|---|---|")
        for a in sc:
            L.append(f"| {a['code']} | {a['label']} | {a.get('level') or '—'} | {fmt(a['points'])}/{fmt(a['max'])} | {a.get('reason','')} |")
    if ns:
        L.append("\n**Não aceitas / sem pontos**\n\n| Código | Área | Justificativa |\n|---|---|---|")
        for a in ns:
            L.append(f"| {a['code']} | {a['label']} | {a.get('reason','')} |")
    if r.get("technical"):
        t = r["technical"]
        parts = [f"{k} {fmt(v[0])}/{fmt(v[1])}" for k, v in t.items() if isinstance(v, list) and len(v) == 2]
        if parts:
            L.append("\n_Análise técnica (2021):_ " + " · ".join(parts))
    if r["regime"] == "NR1_2024":
        L.append("\n**Pontos adicionais (art. 15)**\n\n| Dispositivo | Pedido | Pontos | Decisão | Justificativa |\n|---|---|---|---|---|")
        for b in r.get("bonus", []):
            req = {True: "sim", False: "não", None: "n/d"}.get(b.get("requested"), "n/d")
            dec = "aceito" if float(b["points"]) > 0 else ("não aceito" if b.get("requested") else "não concedido")
            L.append(f"| {b['article']} | {req} | {fmt(b['points'])}/{fmt(b['max'])} | {dec} | {b.get('justification','')} |")
        if not r.get("bonus"):
            L.append("| — | n/d | — | NT sem seção de pontos adicionais | — |")
    elif r.get("bonus"):
        for b in r["bonus"]:
            L.append(f"\n_Ponto adicional:_ {b.get('article')} — {b.get('justification','')}")
    if r.get("notes"):
        L.append(f"\n_Obs.:_ {r['notes']}")
    L.append("")
    return L


def reconcile(recs, xlsx, rowmap):
    """Compare NT area points (+ bonus) with the scoring workbook, item by item."""
    import openpyxl
    rm = json.load(open(rowmap, encoding="utf-8"))
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    proj = list(wb["PROJECTS_ALL"].iter_rows(values_only=True))[1:]
    sc = {str(r[5]).strip(): r for r in list(wb["SCORING_ALL"].iter_rows(values_only=True))[1:]}
    cnt, diffs = Counter(), []
    for aid, r in sorted(recs.items()):
        if r["status"] != "ok" or not r.get("areas_total"):
            cnt[(r["regime"], "sem pontuação por área na NT")] += 1
            continue
        i = rm.get(aid)
        row = sc.get(str(proj[i][6]).strip()) if i is not None else None
        sheet = row[9] if row else None
        if sheet in (None, ""):
            cnt[(r["regime"], "planilha sem nota de áreas/AE")] += 1
            continue
        sheet = float(sheet); at = float(r["areas_total"][0])
        if r["regime"] == "NR1_2024":
            bon = sum(float(b["points"]) for b in r.get("bonus", []))
            calc = min(10.0, at + bon); ok = abs(calc - sheet) <= 0.0015
        else:
            bon, calc = 0.0, at; ok = abs(at - sheet) <= 0.0035
        cnt[(r["regime"], "bate" if ok else "difere")] += 1
        if not ok:
            diffs.append((aid, r["regime"], at, bon, calc, sheet))
    return cnt, diffs


def section3(cnt, diffs):
    L = ["\n## 3. Conciliação das notas das NTs com a planilha (item a item)\n",
         "Para cada avaliação com bloco de áreas na NT, comparou-se a nota de áreas (regime de 2021) ou AE = áreas + pontos adicionais, limitado a 10 "
         "(regime de 2024), com a nota correspondente da planilha de avaliações.\n",
         "| Regime | Bate | Difere | Planilha sem nota | NT sem pontuação por área |", "|---|---|---|---|---|"]
    for reg in ("NR1_2024", "R17_2021"):
        L.append(f"| {reg} | {cnt[(reg,'bate')]} | {cnt[(reg,'difere')]} | {cnt[(reg,'planilha sem nota de áreas/AE')]} | {cnt[(reg,'sem pontuação por área na NT')]} |")
    L.append("\n**Leitura:** as divergências são de dois tipos. (i) A planilha tem mais pontos que a NT, em múltiplos de 0,5 (regime de 2024) ou de 0,133 "
             "(regime de 2021), o que é compatível com a NT ser a análise preliminar e a nota divulgada ser a final, com pontos adicionais ou critérios "
             "reconhecidos depois. (ii) Em poucos casos a planilha tem menos pontos que a NT. A diferença é uma evidência, não uma explicação: nenhuma NT "
             "descreve a mudança.\n")
    L.append("| Avaliação | Regime | Áreas (NT) | Pontos adicionais (NT) | Nota pela NT | Planilha | Diferença |\n|---|---|---|---|---|---|---|")
    for aid, reg, at, bon, calc, sheet in diffs:
        L.append(f"| {aid} | {reg} | {fmt(at)} | {fmt(bon)} | {fmt(calc)} | {fmt(sheet)} | {fmt(sheet - calc)} |")
    return L


def main():
    final, fichas, out = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    xlsx = sys.argv[sys.argv.index("--xlsx") + 1] if "--xlsx" in sys.argv else None
    rowmap = sys.argv[sys.argv.index("--rowmap") + 1] if "--rowmap" in sys.argv else None
    recs = {}
    for f in sorted(final.glob("*.json")):
        recs[f.stem] = json.load(open(f, encoding="utf-8"))
    heads = {p.stem: header(p) for p in (fichas / "avaliacoes").glob("*.md")}
    pheads = {p.stem: header(p) for p in (fichas / "pleitos").glob("*.md")}
    byp = defaultdict(list)
    for aid, r in recs.items():
        byp[heads[aid]["pleito"]].append(aid)
    L = [SECTION1]
    for reg, title in (("NR1_2024", "### 2.1 Regime da Resolução Normativa 1/2024"), ("R17_2021", "### 2.2 Regime da Resolução Cofiex 17/2021")):
        L.append("\n" + title + "\n")
        pids = sorted({p for p, aids in byp.items() if any(recs[a]["regime"] == reg for a in aids)},
                      key=lambda p: (pheads.get(p, {}).get("sector_verified") or pheads.get(p, {}).get("sector", ""), p))
        cur = None
        for p in pids:
            ph = pheads.get(p, {})
            sec = ph.get("sector_verified") or ph.get("sector", "?")
            if sec != cur:
                L.append(f"\n#### Setor: {sec}\n")
                cur = sec
            L.append(f"\n##### Pleito {p} — {ph.get('subtype','')} · {ph.get('lender','')} · faixa {ph.get('ticket_band','')}\n")
            for aid in sorted((a for a in byp[p] if recs[a]["regime"] == reg), key=lambda a: int(heads[a].get("attempt", 0))):
                L += render_eval(recs[aid], heads[aid])
    if xlsx and rowmap:
        cnt, diffs = reconcile(recs, xlsx, rowmap)
        L += section3(cnt, diffs)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(L) + "\n", encoding="utf-8")
    c = Counter((r["regime"], r["status"]) for r in recs.values())
    print(len(recs), "avaliações;", dict(c), "->", out)


if __name__ == "__main__":
    main()
