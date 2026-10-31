"""Add `total_published` (score printed in the official meeting result) to avaliacao fichas whose workbook score is blank.

usage: apply_published_scores.py <manifest.json> <fichas_root> <rowmap.json> <xlsx> <atas_text_dir> [--apply]
Matching is by the 10-digit pleito code printed in the result table; only a single unambiguous score near the code is used.
Never overwrites `total`. Writes a new header key (right after `total`) and `total_published_source: result_<meeting>`.
"""
import glob, json, os, re, sys
from collections import defaultdict
import openpyxl

TOK = re.compile(r"(?<![\d.,])(\d{1,2},\d{3})(?![\d,])")


def main():
    man = json.load(open(sys.argv[1], encoding="utf-8"))
    root = sys.argv[2]
    rowmap = json.load(open(sys.argv[3], encoding="utf-8"))
    wb = openpyxl.load_workbook(sys.argv[4], read_only=True, data_only=True)
    proj = list(wb["PROJECTS_ALL"].iter_rows(values_only=True))[1:]
    sc = {str(r[5]).strip(): r for r in list(wb["SCORING_ALL"].iter_rows(values_only=True))[1:]}
    pub = defaultdict(str)
    for f in glob.glob(sys.argv[5] + "/*.txt"):
        m = re.search(r"Reunião_(\d{3})_", os.path.basename(f))
        if m and ("Resultado" in f or "resultado" in f):
            pub[int(m.group(1))] += open(f, encoding="utf-8", errors="ignore").read() + "\n"
    apply = "--apply" in sys.argv
    done = skipped = 0
    for aid, v in man.items():
        if v["type"] != "avaliacao":
            continue
        fp = f"{root}/avaliacoes/{aid}.md"
        t = open(fp, encoding="utf-8").read()
        if "regime: NR1_2024" not in t or "total_published:" in t:
            continue
        i = rowmap.get(aid)
        r = proj[i]
        row = sc.get(str(r[6]).strip())
        if row and row[9] not in (None, ""):
            continue
        meet = int(re.match(r"\d+", str(r[0])).group())
        txt = man[v["pleito"]]["text"] or ""
        m = re.search(r"Pleito_(\d{8,10})_", txt.replace("\\", "/").split("/")[-1])
        if not m:
            skipped += 1
            continue
        code, p = m.group(1), pub.get(meet, "")
        vals = set()
        for mm in re.finditer(re.escape(code) + r"/20\d\d", p):
            vals |= set(TOK.findall(p[max(0, mm.start() - 150): mm.end() + 450])[:2])
        if len(vals) != 1:
            skipped += 1
            continue
        val = vals.pop().replace(",", ".")
        done += 1
        if apply:
            t2 = re.sub(r"^(total: ?.*)$", rf"\1\ntotal_published: {val}\ntotal_published_source: result_{meet}", t, count=1, flags=re.M)
            open(fp, "w", encoding="utf-8").write(t2)
        print(aid, "meeting", meet, "->", val)
    print("apply" if apply else "dry-run", "| filled:", done, "| skipped (no unambiguous match):", skipped)


if __name__ == "__main__":
    main()
