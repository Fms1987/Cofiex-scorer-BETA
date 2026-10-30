"""Gera references/strategy/sublimit-table.md a partir da base consolidada da COFIEX.

uso: python -I -X utf8 tools/build_sublimit_table.py <COFIEX_Base_Consolidada.xlsx> <skill_dir> [--ate N] [--excluir-pfe X,Y]

--ate N: usa só as reuniões anteriores à N-ésima e trata a N-ésima como a próxima (para
simulações cegas de um pleito julgado na reunião N).
--excluir-pfe: tira da agregação as linhas desses números PFE (o próprio pleito simulado).

Dados agregados por reunião e sublimite: limite previsto, valor aprovado, uso do limite,
contagem de aprovados, não aprovados e retirados, menor nota aprovada e maior nota não aprovada.
Nenhum ente ou projeto é citado. Notas só a partir da 179ª (RN 1/2024); nas reuniões 180ª–184ª
as fontes não trazem a nota dos não aprovados. Rodar de novo a cada reunião.
"""
import json
import os
import sys
from collections import defaultdict

import openpyxl

LIMITE_COL = {  # coluna da aba LIMITES_POR_REUNIAO → categoria da aba PROJETOS
    "Subnacionais – Geral (US$)": "Subnacional – Geral",
    "Subnacionais – PEF e RRF (US$)": "Subnacional – PEF/RRF",
    "Subnacionais – PPPs (US$)": "Subnacional – PPP",
    "Subnacionais – Ambiental e Climático (US$)": "Subnacional – Ambiental e Climático",
    "Subnacionais – Fiscais (US$)": "Subnacional – Fiscal",
    "Subnacionais – Operações em R$ (R$)": "Subnacional – Operações em R$",
    "Federais (US$)": "Federal – Administração Direta",
}
ORDEM = ["Subnacional – Geral", "Subnacional – Ambiental e Climático", "Subnacional – PPP",
         "Subnacional – Fiscal", "Subnacional – Operações em R$", "Subnacional – PEF/RRF",
         "Subnacional – RRF", "Subnacional – Reestruturação de Dívidas",
         "Federal – Administração Direta", "Federal – Empresa Estatal"]
APROVADO = ("Aprovado", "Aprovado parcialmente")


def num(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def nreuniao(r):
    try:
        return int(str(r).split("ª")[0])
    except ValueError:
        return None


def mi(v):
    if v is None:
        return "—"
    return f"{v / 1e6:.1f}".replace(".", ",") if 0 < v < 1e7 else f"{v / 1e6:,.0f}".replace(",", ".")


def nota(v):
    return "—" if v is None else f"{v:.2f}".replace(".", ",")


def main(base, skill, ate=None, excluir=()):
    wb = openpyxl.load_workbook(base, read_only=True, data_only=True)
    lim = list(wb["LIMITES_POR_REUNIAO"].iter_rows(values_only=True))
    hl = {k: i for i, k in enumerate(lim[0])}
    limites, datas, proxima = {}, {}, None
    for r in lim[1:]:
        n = nreuniao(r[0])
        if n is None:
            continue
        datas[n] = r[hl["Data da reunião"]]
        valores = {cat: num(r[hl[col]]) for col, cat in LIMITE_COL.items()}
        if any(v is not None for v in valores.values()):
            limites[n] = valores
            if ate is None and "prevista" in str(r[hl["Data da reunião"]]).lower() or str(r[hl["Formato"]]) == "Prevista":
                proxima = n

    if ate is not None:
        proxima = ate
    proj = list(wb["PROJETOS"].iter_rows(values_only=True))
    hp = {k: i for i, k in enumerate(proj[0])}
    agg = defaultdict(lambda: {"aprov_val": 0.0, "aprov": 0, "nao": 0, "ret": 0, "notas_ap": [], "notas_nao": []})
    for r in proj[1:]:
        n, cat, res = nreuniao(r[hp["Reunião"]]), r[hp["Categoria / Limite"]], r[hp["Resultado"]]
        if (n is None or not cat or (ate is not None and n >= ate)
                or str(r[hp["Nº PFE"]]) in excluir):
            continue
        a = agg[(n, cat)]
        pt = num(r[hp["Pontuação Total"]])
        if res in APROVADO:
            a["aprov"] += 1
            a["aprov_val"] += num(r[hp["Empréstimo aprovado (se parcial)"]] if "Empréstimo aprovado (se parcial)" in hp
                                  and r[hp["Empréstimo aprovado (se parcial)"]] else r[hp["Empréstimo"]]) or 0
            if pt is not None and n >= 179:
                a["notas_ap"].append(pt)
        elif res == "Não aprovado":
            a["nao"] += 1
            if pt is not None and n >= 179:
                a["notas_nao"].append(pt)
        else:
            a["ret"] += 1

    L = ["# Tabela de sublimites por reunião (dados agregados)", "",
         f"Gerada por `tools/build_sublimit_table.py` a partir da base consolidada da COFIEX "
         f"(reuniões {min(datas)}ª–{max(n for n, _ in agg)}ª). Valores em US$ milhões (Operações em R$: R$ milhões). "
         "Nenhum ente ou projeto é citado. Notas = Nota Geral oficial, só sob a RN 1/2024 (179ª em diante); "
         "nas reuniões 180ª–184ª as fontes não trazem a nota dos não aprovados (—).", "",
         "Leitura: os sublimites são anuais e compartilhados entre as reuniões do ano, então o uso por reunião é "
         "indicativo. Uso acima de 100% significa que a reunião aprovou mais que o limite previsto para ela "
         "(saldo de reuniões anteriores ou ajuste).", ""]
    if proxima:
        L += [f"## Próxima reunião: {proxima}ª ({datas[proxima]})", "",
              "| Sublimite | Limite divulgado |", "|---|---|"]
        for cat in ORDEM:
            v = limites[proxima].get(cat)
            if v is not None:
                L.append(f"| {cat} | {mi(v)} |")
        L.append("")
    for cat in ORDEM:
        reunioes = sorted(n for n, c in agg if c == cat)
        if not reunioes:
            continue
        L += [f"## {cat}", "",
              "| Reunião | Data | Limite | Aprovado | Uso do limite | Aprovados | Não aprovados | Retirados de pauta | "
              "Menor nota aprovada | Maior nota não aprovada |",
              "|---|---|---|---|---|---|---|---|---|---|"]
        for n in reunioes:
            a = agg[(n, cat)]
            lv = limites.get(n, {}).get(cat)
            uso = f"{100 * a['aprov_val'] / lv:.0f}%" if lv else ("sem limite próprio" if lv is None else "limite zero")
            L.append(f"| {n}ª | {datas.get(n, '')} | {mi(lv)} | {mi(a['aprov_val'])} | {uso} | {a['aprov']} | {a['nao']} | "
                     f"{a['ret']} | {nota(min(a['notas_ap']) if a['notas_ap'] else None)} | "
                     f"{nota(max(a['notas_nao']) if a['notas_nao'] else None)} |")
        L.append("")
    # Linha de corte recente por sublimite: última reunião (RN 1/2024) com nota de aprovado ou não aprovado.
    cortes = {}
    for cat in ORDEM:
        for n in sorted((n for n, c in agg if c == cat), reverse=True):
            a = agg[(n, cat)]
            if a["notas_ap"] or a["notas_nao"]:
                cortes[cat] = {"reuniao": n,
                               "menor_nota_aprovada": min(a["notas_ap"]) if a["notas_ap"] else None,
                               "maior_nota_nao_aprovada": max(a["notas_nao"]) if a["notas_nao"] else None}
                break
    resumo = {"gerado_de": os.path.basename(base), "ultima_reuniao": max(n for n, _ in agg),
              "proxima_reuniao": {"numero": proxima, "data": str(datas.get(proxima)),
                                  "limites": {c: v for c, v in limites.get(proxima, {}).items() if v is not None}},
              "linha_de_corte_recente": cortes}
    with open(os.path.join(skill, "reference-data", "sublimit_status.json"), "w", encoding="utf-8") as fh:
        json.dump(resumo, fh, ensure_ascii=False, indent=1)
    out = os.path.join(skill, "references", "strategy", "sublimit-table.md")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))
    print(f"gravado {out} ({len(agg)} pares reunião × sublimite)")


if __name__ == "__main__":
    a = sys.argv[1:]
    main(a[0], a[1], int(a[a.index("--ate") + 1]) if "--ate" in a else None,
         a[a.index("--excluir-pfe") + 1].split(",") if "--excluir-pfe" in a else ())
