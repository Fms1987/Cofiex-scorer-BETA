"""Build anonymized calibration cases for cofiex-scorer v17.

Inputs (local only, never committed):
  --xlsx   COFIEX_Avaliacoes_165_185.xlsx  (PROJECTS_ALL, SCORING_ALL)
  --ccroot Cartas-Consultas/               (PDFs + cc_extraidos/*.txt)
Outputs (committed):
  data/cases_anonymized.csv
  skills/cofiex-scorer/references/calibration/cases/<sector>.md  + INDEX.md

Anonymization: proponent, project title, PFE/SEI/CC numbers, meeting number and exact
date are dropped. Opaque ids are assigned from a salted hash whose salt is discarded.
Kept: regime, year+quarter, subtype, lender, ticket band, outcome, reason category,
score components and structural features of the CC text.
"""
import argparse, csv, hashlib, os, re, statistics, subprocess, sys, secrets
from collections import defaultdict
from pathlib import Path
import openpyxl

SECTORS = [  # first match wins; matched against lower-cased title text
    ("water", r"water|sanea|sanit|esgot|hidr|abastec|saneamento|drenag|drainage|flood|inunda|bacia"),
    ("health", r"health|sa[uú]de|hospital|sus\b|vacin|epidem"),
    ("education", r"educa|school|escol|ensino|learning|aprendiz|creche|alfabet"),
    ("justice", r"justi[cç]a|judici|tribunal|court|justice|defensoria|minist[eé]rio p[uú]blico"),
    ("security", r"seguran[cç]a p[uú]blica|public safety|pol[ií]cia|penitenci|prisional|security"),
    ("digital_connectivity", r"conectiv|connectiv|fiber|fibra|telecom|banda larga|broadband|5g|rede de dados"),
    ("digital_govt", r"digital|transforma[cç][aã]o digital|governo eletr|e-gov|inform[aá]tic|tecnologia da informa"),
    ("fiscal", r"fiscal|tribut|pro-?fisco|gest[aã]o p[uú]blica|public management|progest|moderniza[cç][aã]o da gest|reforma|restructur|reestrutura|refinanc|debt"),
    ("environment_climate", r"clima|climate|ambient|environment|biodivers|bioeconom|floresta|forest|carbon|carbono|restaura|ecolog|resili"),
    ("rural", r"rural|agri|agro|agr[ií]col|family farm|semi-?arid|semi[aá]rido|pecu[aá]ri|pesca"),
    ("regional_development", r"territ[oó]r|produtiv|productive|sustainable municip|prosep|competitividade|competitiveness"),
    ("energy", r"energ|el[eé]tric|solar|e[oó]lic|power|transmiss"),
    ("infrastructure_mobility", r"rodov|road|ferrov|metr[oô]|metro|mobilidade|mobility|transport|porto|port\b|aeroport|ponte|bridge|logist|brt|vlt|pavimenta|infraestrutura"),
    ("urban_smart_city", r"urbano|urban|cidade|city|smart|habita|housing|revitaliza|parque|requalifica"),
    ("social", r"social|pobreza|poverty|gender|g[eê]nero|igualdade|racial|inclus|cuidado|juventude|youth|assist"),
    ("fund_intermediation", r"fund|fundo|credit|cr[eé]dito|on-?lending|financ.*access|mpme|msme|bndes|ifi\b|banco"),
]
OUTCOME = {"Approved": "approved", "Approved (partial)": "approved", "Partially approved": "approved",
           "Not approved": "not_approved", "Withdrawn": "withdrawn", "Withdrawn – ARCHIVED": "withdrawn"}
REASONS = [
    (r"limit|budget limit|available", "cap_exceeded"),
    (r"brl|values in", "brl_values"),
    (r"annex i|technical information|incomplete", "annex_i_or_info_gap"),
    (r"capag", "capag_gate"),
    (r"recovery plan|plano de recupera", "fiscal_recovery_plan"),
    (r"art\. ?28", "art28_other"),
    (r"inadequacy|adequa", "budget_inadequacy"),
]
BANDS = [(50, "<50M"), (100, "50-100M"), (200, "100-200M"), (400, "200-400M"), (10**12, ">400M")]


def band(usd, lender="", subtype=""):
    if "BRL" in str(lender) or "BRL" in str(subtype):  # amount column holds reais, not US$
        return "BRL-nominal"
    try:
        m = float(usd) / 1e6
    except (TypeError, ValueError):
        return ""
    return next(lbl for lim, lbl in BANDS if m < lim)


def reason_cat(txt):
    if not txt or txt == "None":
        return ""
    t = str(txt).lower()
    for pat, cat in REASONS:
        if re.search(pat, t):
            return cat
    return "other"


def sector_of(text):
    t = text.lower()
    for name, pat in SECTORS:
        if re.search(pat, t):
            return name
    return "other"


def norm_sei(s):
    return re.sub(r"[^0-9]", "", str(s or ""))[:15]


def quarter(d):
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})", str(d))
    if m:
        return int(m.group(3)), (int(m.group(2)) - 1) // 3 + 1
    m = re.match(r"(\d{4})-(\d{2})", str(d))
    return (int(m.group(1)), (int(m.group(2)) - 1) // 3 + 1) if m else ("", "")


def f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---------- CC text features ----------
def indicators_old(text):
    """2023-24 export: Descricao/Formula/Unidade/Linha de Base/Meta stacked on lines."""
    secs = [m.start() for m in re.finditer(r"2\.6\.\s*Projeto/Indicadores", text, re.I)]
    if not secs:
        return None
    s = secs[-1]
    e = re.search(r"2\.7\.\s*Projeto/Componentes", text[s:], re.I)
    block = text[s: s + (e.start() if e else 6000)]
    nums = re.findall(r"^\s*(-?\d[\d\.]*,\d+|-?\d+)\s*$", block, re.M)
    n = len(nums) // 2
    if n == 0:
        return None
    base = nums[0::2][:n]
    zero = sum(1 for b in base if re.fullmatch(r"-?0+(,0+)?", b))
    return n, 1 - zero / n


def indicators_md_old(text):
    """2023-24 .md export: baseline values follow a lone 'Linha de Base' label as stacked numeric lines."""
    secs = [m.start() for m in re.finditer(r"2\.6\.\s*Projeto/Indicadores", text, re.I)]
    if not secs:
        return None
    block = text[secs[-1]: secs[-1] + 9000]
    m = re.search(r"^\s*Linha de Base\s*$", block, re.M)
    if not m:
        return None
    vals = []
    for line in block[m.end():].splitlines():
        t = line.strip()
        if not t:
            continue
        if re.fullmatch(r"-?\d[\d\.]*(,\d+)?", t):
            vals.append(t)
        else:
            break
    if not vals:
        return None
    # values interleave baselines and targets, and the first target precedes the label: n ~= ceil(len/2).
    # Baseline share is NOT derivable here, so it is left undefined (None); the narrative step verifies it.
    return (len(vals) + 1) // 2, None


def indicators_new(text):
    heads = re.findall(r"Informa[cç][oõ]es do \d+[º°o] indicador", text)
    if not heads:
        return None
    n = len(set(heads))
    bases = re.findall(r"Linha de Base[^:\n]*:\s*([^\n]*)", text)
    if not bases:
        return n, None
    undefined = sum(1 for b in bases if re.match(r"\s*(inexistente|n[aã]o|0+([.,]0+)?\s*$|-|a definir)", b, re.I))
    return n, 1 - undefined / len(bases)


def text_features(text):
    low = text.lower()
    r = indicators_new(text) or indicators_old(text)
    parse = "exact" if r and indicators_new(text) else ("heuristic" if r else "")
    if not r:
        r = indicators_md_old(text)
        parse = "approx" if r else ""
    return {
        "n_indicators_parse": parse,
        "n_indicators": r[0] if r else "",
        "baseline_defined_share": round(r[1], 2) if r and r[1] is not None else "",
        "text_chars_k": round(len(text) / 1000),
        "mentions_opex_sustain": int(bool(re.search(r"sustentabilidade financeira|custos? de (manuten|opera)|opex|operaç[aã]o e manuten", low))),
        "mentions_licensing": int("licenciamento" in low),
        "mentions_gender_race": int(bool(re.search(r"g[eê]nero|racial|ra[cç]a|mulheres", low))),
        "mentions_iso_benchmark": int(bool(re.search(r"iso ?37120|iso ?37122|iso ?37123|\bsnis\b|\bideb\b|\bsaeb\b|cid-?10|datasus|irap", low))),
        "mentions_ppp_guarantee": int(bool(re.search(r"\bppp\b|parceria p[uú]blico|concess", low))),
    }


def _unused_pdf_text(pdf, cache):
    out = cache / (hashlib.md5(str(pdf).encode()).hexdigest() + ".txt")
    if not out.exists():
        subprocess.run(["pdftotext", "-enc", "UTF-8", "-layout", str(pdf), str(out)], check=True)
    return out.read_text(encoding="utf-8", errors="ignore")


def collect_cc(ccroot, cache, pdfroot=None):
    """-> {key: (title_text, features, n_versions)} keyed by PFE number (23-24) or normalized SEI (25-26)."""
    res = {}
    old = Path(pdfroot) / "Pleitos_2023_2024" / "cc_extraidos" if pdfroot else ccroot / "Pleitos_Textos_MD" / "Pleitos_2023_2024"
    byp = defaultdict(list)
    for p in old.rglob("CC_*_Versao_*.txt" if pdfroot else "CC_*_Versao_*.md"):
        m = re.match(r"CC_(\d+)_Versao_(\d+)", p.name)
        byp[m.group(1)].append((int(m.group(2)), p))
    for pfe, vs in byp.items():
        vs.sort()
        text = vs[-1][1].read_text(encoding="utf-8", errors="ignore")
        res[("pfe", pfe)] = (text[:900], text_features(text), len(vs))
    byc = defaultdict(list)
    for p in (ccroot / "Pleitos_Textos_MD" / "Pleitos_2025_2026").rglob("Pleito_*.md"):
        m = re.match(r"Pleito_(\d+)_\d{4}_revisao_(\d+)", p.name)
        if m:
            byc[m.group(1)].append((int(m.group(2)), p))
    for code, vs in byc.items():
        vs.sort()
        text = vs[-1][1].read_text(encoding="utf-8", errors="ignore")
        sei = re.search(r"Processo SEI:\s*([\d\./-]+)", text)
        if sei:
            res[("sei", norm_sei(sei.group(1)))] = (text[:900], text_features(text), len(vs))
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--ccroot", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--cache", required=True)
    ap.add_argument("--pdfroot", default=None, help="folder with Pleitos_2023_2024/cc_extraidos/*.txt")
    a = ap.parse_args()
    out, ccroot, cache = Path(a.out), Path(a.ccroot), Path(a.cache)
    cache.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    projs = list(wb["PROJECTS_ALL"].iter_rows(values_only=True))[1:]
    scores = {}
    for r in list(wb["SCORING_ALL"].iter_rows(values_only=True))[1:]:
        scores[str(r[5]).strip()] = r
    meeting_regime = {}
    for r in scores.values():
        meeting_regime.setdefault(str(r[0]), "NR1_2024" if "1/2024" in str(r[2]) else "R17_2021")
    cc = collect_cc(ccroot, cache, a.pdfroot)

    rows = []
    for r in projs:
        meeting, date, sei, sub, prop, title, pfe, lender, loan, cp, outcome, reason = r[:12]
        pfe = str(pfe).strip()
        c = cc.get(("pfe", pfe)) or cc.get(("sei", norm_sei(sei)))
        sector = sector_of(str(title))
        if sector == "other" and c:  # fall back to the CC text, minus ministry boilerplate
            sector = sector_of(re.sub(r"(?i)assuntos internacionais e desenvolvimento|minist[eé]rio do planejamento", "", c[0]))
        sc = scores.get(pfe)
        regime = meeting_regime.get(str(meeting), "NR1_2024" if str(date)[-4:] >= "2025" else "")
        y, q = quarter(date)
        feats = c[1] if c else {}
        row = {
            "regime": regime, "year": y, "quarter": q,
            "subtype": re.sub(r"[^\x20-\x7e]+", "-", str(sub)).strip("- "),
            "lender": lender, "ticket_band": band(loan, lender, sub),
            "counterpart_pct": round(100 * float(cp) / (float(cp) + float(loan)), 1) if f(cp) is not None and f(loan) else "",
            "sector": sector, "outcome": OUTCOME.get(str(outcome), "other"),
            "reason_cat": reason_cat(reason),
            "capag_pts": "", "debt_pts": "", "ae": "", "idh_pts": "", "total": "",
            "cc_versions": c[2] if c else "",
        }
        if sc and regime == "NR1_2024":
            row.update(capag_pts=sc[6], debt_pts=sc[7], ae=sc[9], idh_pts=sc[10], total=sc[11])
        elif sc:
            row.update(total=sc[11])
        row.update({k: feats.get(k, "") for k in ["n_indicators", "baseline_defined_share", "text_chars_k", "mentions_opex_sustain",
                                                  "mentions_licensing", "mentions_gender_race", "mentions_iso_benchmark", "mentions_ppp_guarantee"]})
        rows.append(row)

    salt = secrets.token_hex(8)  # discarded after the run
    rows.sort(key=lambda r: hashlib.sha256((salt + repr(sorted(r.items()))).encode()).hexdigest())
    for i, r in enumerate(rows, 1):
        r["case_id"] = f"K{i:04d}"
    cols = ["case_id"] + [k for k in rows[0] if k != "case_id"]
    (out / "data").mkdir(parents=True, exist_ok=True)
    with open(out / "data" / "cases_anonymized.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: r["case_id"]))

    # ---- per-sector reference cards ----
    cdir = out / "skills/cofiex-scorer/references/calibration/cases"
    cdir.mkdir(parents=True, exist_ok=True)
    bysec = defaultdict(list)
    for r in rows:
        bysec[r["sector"]].append(r)

    def med(vals):
        vals = [float(v) for v in vals if v not in ("", None)]
        return f"{statistics.median(vals):.2f}" if vals else "n/a"

    index = ["# Anonymized case cards — index", "",
             "Each card = one sector. Contents are anonymized profiles (no proponent, project, number or meeting).",
             "Use them as calibration ranges, **never** as named comparators (see BLIND-JUDGMENT PRINCIPLE).", "",
             "| Sector | n | Approved | Not approved | Withdrawn | AE median (NR 1/2024) | TOTAL median |", "|---|---|---|---|---|---|---|"]
    for sec, rs in sorted(bysec.items(), key=lambda kv: -len(kv[1])):
        n = len(rs)
        cnt = lambda o: sum(1 for r in rs if r["outcome"] == o)
        ae = med([r["ae"] for r in rs if r["regime"] == "NR1_2024"])
        tot = med([r["total"] for r in rs if r["regime"] == "NR1_2024"])
        index.append(f"| [{sec}]({sec}.md) | {n} | {cnt('approved')} | {cnt('not_approved')} | {cnt('withdrawn')} | {ae} | {tot} |")
        lines = [f"# Sector card — {sec}", "",
                 f"n = {n} anonymized cases (meetings 165-185). AE median (NR 1/2024) = {ae}; TOTAL median = {tot}.", "",
                 "| case | regime | yr-Q | subtype | lender | ticket | CP% | outcome | reason | AE | CAPAG | Debt | IDH | TOTAL | #ind | base% | OPEX | ISO/bench |",
                 "|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
        for r in sorted(rs, key=lambda r: r["case_id"]):
            bs = r["baseline_defined_share"]
            lines.append("| {case_id} | {regime} | {year}-Q{quarter} | {subtype} | {lender} | {ticket_band} | {counterpart_pct} | {outcome} | {reason_cat} | {ae} | {capag_pts} | {debt_pts} | {idh_pts} | {total} | {n_indicators} | {bs} | {o} | {i} |".format(
                bs=(f"{int(float(bs)*100)}" if bs != "" else ""), o=r["mentions_opex_sustain"], i=r["mentions_iso_benchmark"], **r))
        (cdir / f"{sec}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    (cdir / "INDEX.md").write_text("\n".join(index) + "\n", encoding="utf-8")
    print(len(rows), "cases;", {k: len(v) for k, v in bysec.items()})


if __name__ == "__main__":
    main()
