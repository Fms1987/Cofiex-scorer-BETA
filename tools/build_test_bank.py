"""Freeze the fixed test bank for regression runs (proposal item C6).

usage: build_test_bank.py <match_csv> <consolidated_xlsx> <idmap_json> <out_dir> [--pool ids.csv] [--n 40] [--seed 2026]

match_csv: output of match_pleitos_to_results.py (pleito_id, arquivo, sei, reuniao, n_pfe, ...).
idmap_json: the PRIVATE map "P|<pfe or sei digits>" -> opaque ficha id (kept outside the repository).
--pool: optional CSV with a pleito_id column (and nota_oficial, to pick the meeting that was scored);
        restricts candidates (e.g. pleitos already run blind, so the previous results become the baseline).

Writes to <out_dir> (PRIVATE: holds official scores; never copy it into the skill):
  bank.csv         pleito_id, opaque_id, arquivo, reuniao, estrato
  ids.txt          opaque ids, one per line (input of make_blind_sandbox.py @ids.txt)
  truth.csv        official Nota Geral and components per pleito
  truth_areas.csv  official level per PPA objective (A-E, 0 = pleiteada sem nota)
Only NR 1/2024 pleitos with an official total and an opaque id are eligible. Sampling is stratified
by sublimit group (Geral / other subnational / Federal) x entity type, proportional, min 1 per stratum.
"""
import argparse, csv, json, random, re
from collections import defaultdict
from pathlib import Path

import openpyxl

PONTOS_NIVEL = {5.0: "A", 3.0: "B", 2.0: "C", 1.0: "D", 0.5: "E", 0.0: "0"}


def pfe_key(v):
    return str(v or "").strip()


def num(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def grupo(cat):
    cat = cat or ""
    if cat.startswith("Federal"):
        return "Federal"
    return "Geral" if "Geral" in cat else "Subnacional-especifico"


def opaque(idmap, row):
    for k in ("P|" + re.sub(r"\D", "", row["sei"].split("-")[0]), "P|" + row["n_pfe"].split("/")[0]):
        if k in idmap:
            return idmap[k]
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("match_csv"); ap.add_argument("xlsx"); ap.add_argument("idmap"); ap.add_argument("out_dir")
    ap.add_argument("--pool"); ap.add_argument("--n", type=int, default=40); ap.add_argument("--seed", type=int, default=2026)
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    pont, cat = {}, {}
    for r in wb["PONTUACAO"].iter_rows(min_row=2, values_only=True):
        if r[0] and "1/2024" in str(r[2] or ""):
            pont[(pfe_key(r[3]), r[0])] = {"capag": num(r[8]), "divida": num(r[9]), "ae": num(r[17]), "idh": num(r[18]),
                                          "nota_geral": num(r[20]), "resultado": r[7], "tipo_ente": r[4]}
    for r in wb["PROJETOS"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            cat[(pfe_key(r[2]), r[0])] = r[8]
    areas = defaultdict(list)
    for r in wb["NT_AREAS_ESTRATEGICAS"].iter_rows(min_row=2, values_only=True):
        if r[0] and "1/2024" in str(r[4] or "") and r[5]:
            pts = num(r[8]) or 0.0
            areas[pfe_key(r[1])].append({"reuniao": r[0], "codigo_ppa": str(r[5]), "nivel": r[7] or PONTOS_NIVEL.get(pts, "?"), "pontos": pts})

    idmap = json.loads(Path(a.idmap).read_text(encoding="utf-8"))
    pool = None
    if a.pool:
        with open(a.pool, encoding="utf-8") as fh:
            pool = {r["pleito_id"]: num(r.get("nota_oficial")) for r in csv.DictReader(fh)}
    cands = {}
    with open(a.match_csv, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if pool is not None and r["pleito_id"] not in pool:
                continue
            t = pont.get((r["n_pfe"], r["reuniao"]))
            oid = opaque(idmap, r)
            if not t or t["nota_geral"] is None or not oid:
                continue
            r.update(t, opaque_id=oid, categoria=cat.get((r["n_pfe"], r["reuniao"]), ""))
            r["estrato"] = f"{grupo(r['categoria'])}|{t['tipo_ente']}"
            # one row per pleito: the meeting whose total the pool scored against, else the latest meeting
            alvo = pool.get(r["pleito_id"]) if pool else None
            prev = cands.get(r["pleito_id"])
            if alvo is not None:
                if abs(t["nota_geral"] - alvo) < 0.001:
                    cands[r["pleito_id"]] = r
            elif not prev or int(re.sub(r"\D", "", r["reuniao"])) > int(re.sub(r"\D", "", prev["reuniao"])):
                cands[r["pleito_id"]] = r
    cands = list(cands.values())

    by = defaultdict(list)
    for r in cands:
        by[r["estrato"]].append(r)
    n = min(a.n, len(cands))
    quota = {s: max(1, n * len(v) // len(cands)) for s, v in by.items()}
    rest = sorted(by, key=lambda s: -(n * len(by[s]) / len(cands) - n * len(by[s]) // len(cands)))
    i = 0
    while sum(quota.values()) < n:
        s = rest[i % len(rest)]
        if quota[s] < len(by[s]):
            quota[s] += 1
        i += 1
    rng = random.Random(a.seed)
    bank = []
    for s in sorted(by):
        bank += rng.sample(sorted(by[s], key=lambda r: r["pleito_id"]), min(quota[s], len(by[s])))

    out = Path(a.out_dir); out.mkdir(parents=True, exist_ok=True)
    with open(out / "bank.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(["pleito_id", "opaque_id", "arquivo", "reuniao", "estrato"])
        for r in bank:
            w.writerow([r["pleito_id"], r["opaque_id"], r["arquivo"], r["reuniao"], r["estrato"]])
    (out / "ids.txt").write_text("\n".join(r["opaque_id"] for r in bank) + "\n", encoding="utf-8")
    with open(out / "truth.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(["pleito_id", "categoria", "tipo_ente", "resultado", "nota_geral", "ae", "capag", "divida", "idh"])
        for r in bank:
            w.writerow([r["pleito_id"], r["categoria"], r["tipo_ente"], r["resultado"], r["nota_geral"], r["ae"], r["capag"], r["divida"], r["idh"]])
    n_areas = 0
    with open(out / "truth_areas.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(["pleito_id", "codigo_ppa", "nivel", "pontos"])
        for r in bank:
            rows = areas.get(r["n_pfe"], [])
            same = [x for x in rows if x["reuniao"] == r["reuniao"]]
            if not same and rows:
                last = max(x["reuniao"] for x in rows)
                same = [x for x in rows if x["reuniao"] == last]
            for x in same:
                w.writerow([r["pleito_id"], x["codigo_ppa"], x["nivel"], x["pontos"]]); n_areas += 1
    print(f"{len(cands)} candidatos; banco com {len(bank)} pleitos; {n_areas} linhas de área")
    for s in sorted(by):
        print(f"  {s}: {quota[s]} de {len(by[s])}")


if __name__ == "__main__":
    main()
