"""Find pleitos that have an official score (PONTUACAO sheet) in the consolidated workbook but
no matching NT file (NT_PONTUACAO sheet, "Arquivo da NT" column empty or row missing), then check
whether a raw carta-consulta .md text exists for them (Pleitos_Textos_MD, both 2023_2024 "CC_<pfe>"
and 2025_2026 "Pleito_<pfe>_..." naming) so they can be used as new blind-test cases.

Usage: python tools/find_pleitos_sem_nt.py <xlsx> <pleitos_2023_2024_dir> <pleitos_2025_2026_dir> <out_csv>
"""
import argparse
import csv
import re
from pathlib import Path

import openpyxl


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("dir_2023_2024")
    ap.add_argument("dir_2025_2026")
    ap.add_argument("out_csv")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    pont_rows = list(wb["PONTUACAO"].iter_rows(values_only=True))[1:]
    nt_rows = list(wb["NT_PONTUACAO"].iter_rows(values_only=True))[1:]

    has_nt = set()
    for r in nt_rows:
        reuniao, pfe, arquivo_nt = r[0], r[1], r[18]
        if arquivo_nt:
            has_nt.add((reuniao, str(pfe)))

    with_nota = {}
    for r in pont_rows:
        reuniao, pfe, resultado, nota_final = r[0], r[3], r[7], r[20]
        if nota_final is not None:
            with_nota[(reuniao, str(pfe))] = {
                "reuniao": reuniao, "n_pfe": pfe, "tipo_ente": r[4], "nome_ente": r[5],
                "nome_projeto": r[6], "resultado": resultado, "nota_final": nota_final,
            }

    sem_nt = {k: v for k, v in with_nota.items() if k not in has_nt}
    print(f"{len(with_nota)} pleitos com nota final; {len(sem_nt)} sem NT (Arquivo da NT vazio/ausente).")

    # Index available carta-consulta text files by PFE number. 2023_2024 pleitos are stored as
    # CC_<pfe>/CC_<pfe>_Versao_<n>.md; keep only the latest version per pleito.
    text_by_pfe = {}
    for d in Path(a.dir_2023_2024).glob("CC_*"):
        if not d.is_dir():
            continue
        m = re.match(r"CC_(\d+)", d.name)
        if not m:
            continue
        pfe = m.group(1)
        versions = []
        for f in d.glob("CC_*_Versao_*.md"):
            vm = re.search(r"Versao_(\d+)", f.stem)
            if vm:
                versions.append((int(vm.group(1)), f))
        if versions:
            versions.sort()
            text_by_pfe[pfe] = [versions[-1][1]]
    rev_candidates = {}
    for f in Path(a.dir_2025_2026).glob("Pleito_*.md"):
        m = re.match(r"Pleito_(\d+)_(\d+)_revisao_(\d+)_(\S+)", f.stem)
        if not m:
            continue
        pid, _, rev, ts = m.groups()
        key = (int(rev), ts)
        if pid not in rev_candidates or key > rev_candidates[pid][0]:
            rev_candidates[pid] = (key, f)
    for pid, (_, f) in rev_candidates.items():
        text_by_pfe[pid] = [f]

    out_rows = []
    for (reuniao, pfe), v in sorted(sem_nt.items()):
        pfe_key = str(pfe).split("/")[0]
        files = text_by_pfe.get(pfe_key, [])
        out_rows.append({
            **v, "tem_texto_carta": "SIM" if files else "NAO",
            "arquivo_carta": files[-1].name if files else "",
        })

    with open(a.out_csv, "w", newline="", encoding="utf-8") as fh:
        fieldnames = ["reuniao", "n_pfe", "tipo_ente", "nome_ente", "nome_projeto", "resultado",
                      "nota_final", "tem_texto_carta", "arquivo_carta"]
        w = csv.DictWriter(fh, fieldnames=fieldnames)
        w.writeheader()
        for r in out_rows:
            w.writerow(r)

    com_texto = sum(1 for r in out_rows if r["tem_texto_carta"] == "SIM")
    print(f"{com_texto}/{len(out_rows)} pleitos sem NT têm texto de carta-consulta disponível.")
    print(f"Saída: {a.out_csv}")


if __name__ == "__main__":
    main()
